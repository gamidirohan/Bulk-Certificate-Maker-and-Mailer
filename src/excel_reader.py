import re
from email.utils import parseaddr

from openpyxl import load_workbook


def _normalized(value):
    return re.sub(r"[^a-z]", "", str(value or "").lower())


def read_participants(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Workbook is empty")
    headers = [_normalized(value) for value in rows[0]]
    name_index = next((i for i, value in enumerate(headers) if value in {"name", "fullname", "participantname"} or "name" in value), None)
    email_index = next((i for i, value in enumerate(headers) if value in {"email", "emailaddress"} or "email" in value), None)
    if name_index is None or email_index is None:
        raise ValueError("Could not identify Full Name and Email Address columns.")
    participants = []
    for row_number, row in enumerate(rows[1:], 2):
        name = str(row[name_index] or "").strip()
        email = str(row[email_index] or "").strip()
        valid = bool(parseaddr(email)[1]) and "@" in email and "." in email.split("@")[-1]
        if name or email:
            participants.append({"row": row_number, "name": name, "email": email if valid else ""})
    return participants
